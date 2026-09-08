"""Turn an uploaded file into `(page_number, text)` pairs.

Page numbers are 1-based and only meaningful for PDFs; other formats yield a
single `None` page. Downstream, that becomes the page shown in a citation.
"""

from __future__ import annotations

import io
import logging
from pathlib import Path

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS: frozenset[str] = frozenset(
    {".pdf", ".docx", ".txt", ".md", ".markdown", ".rst", ".csv", ".json"}
)

# Browsers are inconsistent about the MIME type they attach to an upload, so
# the extension is the authority and this map is only used for display.
EXTENSION_MIME: dict[str, str] = {
    ".pdf": "application/pdf",
    ".docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".markdown": "text/markdown",
    ".rst": "text/x-rst",
    ".csv": "text/csv",
    ".json": "application/json",
}


class UnsupportedFileError(ValueError):
    """The file extension is not one we can extract text from."""


class EmptyDocumentError(ValueError):
    """The file parsed cleanly but contained no extractable text."""


def is_supported(filename: str) -> bool:
    return Path(filename).suffix.lower() in SUPPORTED_EXTENSIONS


def guess_mime(filename: str) -> str:
    return EXTENSION_MIME.get(Path(filename).suffix.lower(), "application/octet-stream")


def load_pages(filename: str, data: bytes) -> list[tuple[int | None, str]]:
    """Extract text from `data`, dispatching on the extension of `filename`.

    Raises `UnsupportedFileError` for unknown types and `EmptyDocumentError`
    when nothing readable came out — the latter is common with scanned PDFs,
    which need OCR this project does not attempt.
    """
    suffix = Path(filename).suffix.lower()

    if suffix == ".pdf":
        pages = _load_pdf(data)
    elif suffix == ".docx":
        pages = _load_docx(data)
    elif suffix in SUPPORTED_EXTENSIONS:
        pages = [(None, _decode(data))]
    else:
        raise UnsupportedFileError(
            f"Cannot read '{filename}'. Supported types: "
            f"{', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    pages = [(page, text) for page, text in pages if text and text.strip()]
    if not pages:
        raise EmptyDocumentError(
            f"No text could be extracted from '{filename}'. If this is a scanned "
            f"PDF it contains images rather than text, and would need OCR first."
        )
    return pages


def _load_pdf(data: bytes) -> list[tuple[int | None, str]]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages: list[tuple[int | None, str]] = []
    for number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:  # noqa: BLE001 - one broken page must not kill the upload
            logger.warning("Could not extract text from PDF page %d", number)
            continue
        pages.append((number, _normalize(text)))
    return pages


def _load_docx(data: bytes) -> list[tuple[int | None, str]]:
    import docx

    document = docx.Document(io.BytesIO(data))

    blocks = [p.text for p in document.paragraphs if p.text and p.text.strip()]
    # Tables often carry the substance in specs and reports; losing them silently
    # is a common cause of "the answer isn't in my documents".
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))

    return [(None, _normalize("\n\n".join(blocks)))]


def _decode(data: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return _normalize(data.decode(encoding))
        except UnicodeDecodeError:
            continue
    return _normalize(data.decode("utf-8", errors="replace"))


def _normalize(text: str) -> str:
    """Normalise line endings and collapse the runs of blank lines PDFs produce."""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    lines = [line.rstrip() for line in text.split("\n")]

    out: list[str] = []
    blanks = 0
    for line in lines:
        if line:
            blanks = 0
            out.append(line)
        else:
            blanks += 1
            if blanks <= 1:
                out.append("")
    return "\n".join(out).strip()
